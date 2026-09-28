"""Approved regional entity review, deduplication, and export artifacts."""

from __future__ import annotations

import re
from collections import defaultdict
from copy import copy
from hashlib import sha256
from io import BytesIO
from typing import Any


def build_regional_review(
    region: str,
    model_tree: list[dict[str, Any]],
    proposals: dict[str, dict[str, Any]],
    decisions: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    """Build a lossless approved regional view and its source-to-output truth map."""
    sources: list[dict[str, Any]] = []
    for api_branch in model_tree:
        for model in api_branch["models"]:
            source_key = f"{api_branch['runId']}:{model['id']}"
            proposal = proposals.get(source_key, {})
            decision = decisions.get(source_key, {})
            use_ai = decision.get("selection") == "AI suggestion" and bool(proposal)
            proposed_attributes = {
                item["attributeId"]: item for item in proposal.get("attributes", [])
            }
            attributes = []
            for field in model["fields"]:
                field_decision = decision.get("attributes", {}).get(field["id"], {})
                suggested = proposed_attributes.get(field["id"], {})
                use_field_ai = field_decision.get("selection") == "AI suggestion" and bool(
                    suggested
                )
                attributes.append(
                    {
                        "sourceAttributeId": field["id"],
                        "originalName": field["Field"],
                        "approvedName": (
                            suggested.get("normalizedName", field["Field"])
                            if use_field_ai
                            else field["Field"]
                        ),
                        "approvedDescription": (
                            suggested.get("normalizedDescription", field["Description"])
                            if use_field_ai
                            else field["Description"]
                        ),
                        "type": field["Type"],
                        "required": field["Required"],
                        "selection": "AI suggestion" if use_field_ai else "Original",
                        "comment": field_decision.get("comment", "").strip(),
                    }
                )
            sources.append(
                {
                    "sourceKey": source_key,
                    "runId": api_branch["runId"],
                    "api": api_branch["api"],
                    "repository": api_branch["repository"],
                    "sourceEntityId": model["id"],
                    "originalName": model["name"],
                    "approvedName": (
                        proposal.get("normalizedName", model["name"]) if use_ai else model["name"]
                    ),
                    "approvedDescription": (
                        proposal.get("normalizedDescription", model["description"])
                        if use_ai
                        else model["description"]
                    ),
                    "selection": "AI suggestion" if use_ai else "Original",
                    "comment": decision.get("comment", "").strip(),
                    "endpoints": model["mappings"],
                    "attributes": attributes,
                }
            )

    entity_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for source in sources:
        entity_groups[_comparison_key(source["approvedName"])].append(source)

    regional_entities = []
    truth_mappings = []
    removed_duplicates = []
    for entity_key, members in sorted(entity_groups.items()):
        members.sort(key=lambda item: (item["api"].casefold(), item["sourceKey"]))
        retained = members[0]
        regional_entity_id = _stable_id("regional-entity", region, entity_key)
        for index, member in enumerate(members):
            entity_mapping = _truth_mapping(
                region,
                regional_entity_id,
                retained["approvedName"],
                None,
                None,
                member,
                None,
                "Retained" if index == 0 else "Merged duplicate entity",
            )
            truth_mappings.append(entity_mapping)
            if index > 0:
                removed_duplicates.append(entity_mapping)
        attribute_groups: dict[tuple[str, str], list[tuple[dict[str, Any], dict[str, Any]]]] = (
            defaultdict(list)
        )
        for member in members:
            for attribute in member["attributes"]:
                attribute_groups[
                    (_comparison_key(attribute["approvedName"]), attribute["type"].casefold())
                ].append((member, attribute))

        regional_attributes = []
        for (attribute_key, type_key), occurrences in sorted(attribute_groups.items()):
            occurrences.sort(
                key=lambda pair: (pair[0]["api"].casefold(), pair[1]["sourceAttributeId"])
            )
            retained_member, retained_attribute = occurrences[0]
            regional_attribute_id = _stable_id(
                "regional-attribute", regional_entity_id, attribute_key, type_key
            )
            regional_attributes.append(
                {
                    "id": regional_attribute_id,
                    "name": retained_attribute["approvedName"],
                    "description": retained_attribute["approvedDescription"],
                    "type": retained_attribute["type"],
                    "requiredInAllSources": all(item[1]["required"] for item in occurrences),
                    "sourceCount": len(occurrences),
                }
            )
            for index, (member, attribute) in enumerate(occurrences):
                mapping = _truth_mapping(
                    region,
                    regional_entity_id,
                    retained["approvedName"],
                    regional_attribute_id,
                    retained_attribute["approvedName"],
                    member,
                    attribute,
                    "Retained" if index == 0 else "Merged duplicate attribute",
                )
                truth_mappings.append(mapping)
                if index > 0:
                    removed_duplicates.append(mapping)

        regional_entities.append(
            {
                "id": regional_entity_id,
                "name": retained["approvedName"],
                "description": retained["approvedDescription"],
                "attributes": regional_attributes,
                "sourceEntityCount": len(members),
                "duplicateEntityCount": max(0, len(members) - 1),
                "apis": sorted({member["api"] for member in members}),
                "endpoints": _unique_endpoints(members),
            }
        )

    return {
        "version": "1.0",
        "region": region,
        "status": "Approved",
        "deduplicationRule": (
            "Entities merge by reviewer-approved name; attributes merge within that entity by "
            "reviewer-approved name and source type. Source artifacts are unchanged."
        ),
        "summary": {
            "sourceEntities": len(sources),
            "regionalEntities": len(regional_entities),
            "mergedDuplicateEntities": sum(
                entity["duplicateEntityCount"] for entity in regional_entities
            ),
            "sourceAttributes": sum(len(source["attributes"]) for source in sources),
            "regionalAttributes": sum(len(entity["attributes"]) for entity in regional_entities),
        },
        "entities": regional_entities,
        "truthMappings": truth_mappings,
        "removedDuplicates": removed_duplicates,
        "sourceReviews": sources,
    }


def render_regional_review_excel(review: dict[str, Any]) -> bytes:
    """Render the approved regional model and truth mapping as a styled XLSX workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    entities = workbook.create_sheet("Regional entities")
    truth = workbook.create_sheet("Truth mapping")
    duplicates = workbook.create_sheet("Removed duplicates")

    summary.append([f"Approved regional entity review — {review['region']}"])
    summary.append([])
    summary.append(["Metric", "Value"])
    for label, value in review["summary"].items():
        summary.append([_title(label), value])
    summary.append([])
    summary.append(["Deduplication rule", review["deduplicationRule"]])

    entity_headers = [
        "Regional entity ID",
        "Regional entity",
        "Description",
        "Regional attribute ID",
        "Attribute",
        "Type",
        "Required in all sources",
        "Source count",
        "APIs",
        "Endpoints",
    ]
    entities.append(entity_headers)
    for entity in review["entities"]:
        endpoint_text = "; ".join(
            f"{item['API']} {item['Method']} {item['Route']} ({item['Usage']})"
            for item in entity["endpoints"]
        )
        for attribute in entity["attributes"] or [{}]:
            entities.append(
                [
                    entity["id"],
                    entity["name"],
                    entity["description"],
                    attribute.get("id", ""),
                    attribute.get("name", ""),
                    attribute.get("type", ""),
                    attribute.get("requiredInAllSources", ""),
                    attribute.get("sourceCount", 0),
                    ", ".join(entity["apis"]),
                    endpoint_text,
                ]
            )

    truth_headers = [
        "Action",
        "API",
        "Repository",
        "Source entity ID",
        "Source entity",
        "Source attribute ID",
        "Source attribute",
        "Regional entity ID",
        "Regional entity",
        "Regional attribute ID",
        "Regional attribute",
        "Endpoints involved",
        "Reviewer selection",
        "Reviewer comment",
    ]
    truth.append(truth_headers)
    duplicates.append(truth_headers)
    for mapping in review["truthMappings"]:
        truth.append(_truth_row(mapping))
    for mapping in review["removedDuplicates"]:
        duplicates.append(_truth_row(mapping))

    dark_fill = PatternFill("solid", fgColor="1F4E78")
    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2" if sheet.title != "Summary" else "A4"
        header_row = 3 if sheet.title == "Summary" else 1
        for cell in sheet[header_row]:
            cell.fill = dark_fill
            cell.font = Font(color="FFFFFF", bold=True, name="Arial", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = copy(cell.font)
                cell.font = Font(
                    name="Arial",
                    size=10,
                    bold=cell.font.bold,
                    italic=cell.font.italic,
                    color=cell.font.color,
                )
                cell.alignment = Alignment(vertical="top", wrap_text=False)
        for column_cells in sheet.columns:
            letter = column_cells[0].column_letter
            width = min(48, max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2))
            sheet.column_dimensions[letter].width = width
        if sheet.title == "Summary":
            sheet["A1"].font = Font(name="Arial", size=14, bold=True, color="1F1F1F")
            sheet.column_dimensions["A"].width = 32
            sheet.column_dimensions["B"].width = 88

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_regional_review_mermaid(review: dict[str, Any]) -> bytes:
    """Render source, endpoint, approved model, and duplicate lineage as Mermaid."""
    lines = ["flowchart LR", f'  region["Region: {_mermaid(review["region"])}"]']
    seen: set[str] = {"region"}
    for entity in review["entities"]:
        entity_node = _node_id(entity["id"])
        lines.append(f'  {entity_node}["{_mermaid(entity["name"])}"]')
        lines.append(f"  region --> {entity_node}")
        for attribute in entity["attributes"]:
            attribute_node = _node_id(attribute["id"])
            attribute_label = f"{_mermaid(attribute['name'])}: {_mermaid(attribute['type'])}"
            lines.append(f'  {attribute_node}["{attribute_label}"]')
            lines.append(f"  {entity_node} --> {attribute_node}")
    for mapping in review["truthMappings"]:
        api_node = _node_id("api-" + mapping["api"])
        source_node = _node_id("source-" + mapping["sourceEntityId"] + mapping["api"])
        if api_node not in seen:
            lines.append(f'  {api_node}[["API: {_mermaid(mapping["api"])}"]]')
            seen.add(api_node)
        if source_node not in seen:
            lines.append(f'  {source_node}["Source: {_mermaid(mapping["sourceEntity"])}"]')
            lines.append(f"  {api_node} --> {source_node}")
            edge = "-. merged .->" if mapping["action"] != "Retained" else "-->"
            lines.append(f"  {source_node} {edge} {_node_id(mapping['regionalEntityId'])}")
            seen.add(source_node)
        for endpoint in mapping["endpoints"]:
            endpoint_node = _node_id("endpoint-" + endpoint)
            if endpoint_node not in seen:
                lines.append(f'  {endpoint_node}(["{_mermaid(endpoint)}"])')
                lines.append(f"  {endpoint_node} --> {source_node}")
                seen.add(endpoint_node)
        if mapping["sourceAttributeId"]:
            source_attribute_node = _node_id(
                "source-attribute-" + mapping["sourceAttributeId"] + mapping["api"]
            )
            if source_attribute_node not in seen:
                lines.append(
                    f'  {source_attribute_node}["Source field: '
                    f'{_mermaid(mapping["sourceAttribute"])}"]'
                )
                lines.append(f"  {source_node} --> {source_attribute_node}")
                attribute_edge = (
                    "-. merged .->" if mapping["action"] == "Merged duplicate attribute" else "-->"
                )
                lines.append(
                    f"  {source_attribute_node} {attribute_edge} "
                    f"{_node_id(mapping['regionalAttributeId'])}"
                )
                seen.add(source_attribute_node)
    return ("\n".join(lines) + "\n").encode()


def _truth_mapping(
    region: str,
    regional_entity_id: str,
    regional_entity: str,
    regional_attribute_id: str | None,
    regional_attribute: str | None,
    member: dict[str, Any],
    attribute: dict[str, Any] | None,
    action: str,
) -> dict[str, Any]:
    return {
        "region": region,
        "action": action,
        "api": member["api"],
        "repository": member["repository"],
        "sourceEntityId": member["sourceEntityId"],
        "sourceEntity": member["originalName"],
        "sourceAttributeId": attribute["sourceAttributeId"] if attribute else "",
        "sourceAttribute": attribute["originalName"] if attribute else "",
        "regionalEntityId": regional_entity_id,
        "regionalEntity": regional_entity,
        "regionalAttributeId": regional_attribute_id or "",
        "regionalAttribute": regional_attribute or "",
        "endpoints": [
            f"{member['api']} {item['Endpoint']} · {item['Method']} {item['Route']} "
            f"({item['Usage']})"
            for item in member["endpoints"]
        ],
        "reviewerSelection": attribute["selection"] if attribute else member["selection"],
        "reviewerComment": attribute["comment"] if attribute else member["comment"],
    }


def _truth_columns() -> list[str]:
    return [
        "action",
        "api",
        "repository",
        "sourceEntityId",
        "sourceEntity",
        "sourceAttributeId",
        "sourceAttribute",
        "regionalEntityId",
        "regionalEntity",
        "regionalAttributeId",
        "regionalAttribute",
        "endpoints",
        "reviewerSelection",
        "reviewerComment",
    ]


def _truth_row(mapping: dict[str, Any]) -> list[Any]:
    return [
        "; ".join(mapping.get(key, [])) if key == "endpoints" else mapping.get(key, "")
        for key in _truth_columns()
    ]


def _unique_endpoints(members: list[dict[str, Any]]) -> list[dict[str, str]]:
    unique = {
        (member["api"], item["Method"], item["Route"], item["Usage"]): {
            "API": member["api"],
            **item,
        }
        for member in members
        for item in member["endpoints"]
    }
    return [unique[key] for key in sorted(unique)]


def _comparison_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def _stable_id(kind: str, *values: str) -> str:
    digest = sha256("\x1f".join(values).encode()).hexdigest()[:20]
    return f"{kind}-{digest}"


def _node_id(value: str) -> str:
    return "n" + sha256(value.encode()).hexdigest()[:12]


def _mermaid(value: Any) -> str:
    return str(value).replace('"', "'").replace("\n", " ")


def _title(value: str) -> str:
    return re.sub(r"(?<!^)(?=[A-Z])", " ", value).capitalize()
