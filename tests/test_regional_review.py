from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from canonical_model_generator.regional_review import (
    build_regional_review,
    render_regional_review_excel,
    render_regional_review_mermaid,
)


def _tree() -> list[dict]:
    mappings = [
        {
            "Endpoint": "Create quote",
            "Method": "POST",
            "Route": "/quotes",
            "Usage": "Request",
        }
    ]
    return [
        {
            "runId": run_id,
            "api": api,
            "repository": f"{api}.zip",
            "models": [
                {
                    "id": f"entity-{run_id}",
                    "name": source_name,
                    "description": f"Source {source_name}",
                    "fields": [
                        {
                            "id": f"attribute-{run_id}",
                            "Field": field_name,
                            "Type": "string",
                            "Required": True,
                            "Description": f"Source {field_name}",
                        }
                    ],
                    "mappings": mappings,
                }
            ],
        }
        for run_id, api, source_name, field_name in (
            ("one", "quote-api", "QuoteRequest", "postalCode"),
            ("two", "pricing-api", "PricingInput", "zipCode"),
        )
    ]


def test_review_merges_approved_names_and_preserves_truth() -> None:
    proposals = {
        "one:entity-one": {
            "normalizedName": "RegionalQuote",
            "normalizedDescription": "Approved quote input.",
            "attributes": [
                {
                    "attributeId": "attribute-one",
                    "normalizedName": "postalCode",
                    "normalizedDescription": "Approved postal code.",
                }
            ],
        },
        "two:entity-two": {
            "normalizedName": "RegionalQuote",
            "normalizedDescription": "Approved quote input.",
            "attributes": [
                {
                    "attributeId": "attribute-two",
                    "normalizedName": "postalCode",
                    "normalizedDescription": "Approved postal code.",
                }
            ],
        },
    }
    decisions = {
        key: {
            "selection": "AI suggestion",
            "comment": "Approved entity",
            "attributes": {
                f"attribute-{suffix}": {
                    "selection": "AI suggestion",
                    "comment": "Approved field",
                }
            },
        }
        for key, suffix in (("one:entity-one", "one"), ("two:entity-two", "two"))
    }

    review = build_regional_review("EU", _tree(), proposals, decisions)

    assert review["summary"] == {
        "sourceEntities": 2,
        "regionalEntities": 1,
        "mergedDuplicateEntities": 1,
        "sourceAttributes": 2,
        "regionalAttributes": 1,
    }
    assert {item["api"] for item in review["truthMappings"]} == {
        "quote-api",
        "pricing-api",
    }
    assert sum(not item["sourceAttributeId"] for item in review["truthMappings"]) == 2
    assert {item["action"] for item in review["removedDuplicates"]} == {
        "Merged duplicate entity",
        "Merged duplicate attribute",
    }
    assert all(item["endpoints"] for item in review["removedDuplicates"])


def test_review_exports_openable_excel_and_mermaid_lineage() -> None:
    review = build_regional_review("EU", _tree(), {}, {})

    workbook = load_workbook(BytesIO(render_regional_review_excel(review)))
    mermaid = render_regional_review_mermaid(review).decode()

    assert workbook.sheetnames == [
        "Summary",
        "Regional entities",
        "Truth mapping",
        "Removed duplicates",
    ]
    assert workbook["Truth mapping"].max_row > 1
    assert mermaid.startswith("flowchart LR")
    assert "quote-api" in mermaid
    assert "Create quote" in mermaid
